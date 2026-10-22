#!/usr/bin/env python3
"""So từng câu trong file .docx gốc với trang HTML đã build để chắc không sót đoạn nào.

Chạy sau `npm run build`:  python3 -I scripts/check_fidelity.py [--show]
- Câu ngắn (< 25 ký tự sau chuẩn hoá) được bỏ qua.
- So khớp "mềm": bỏ dấu câu, ngoặc kép, khoảng trắng; vẫn giữ dấu tiếng Việt.
- Câu bị sửa lỗi chính tả có thể báo "thiếu" — đối chiếu với docs/ghi-chu-bien-tap/.
"""
import html
import re
import sys
import unicodedata
from pathlib import Path

import docx

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / 'scripts'))
from convert_docx import SLUGS  # noqa: E402

SHOW = '--show' in sys.argv


def norm(s: str) -> str:
    s = unicodedata.normalize('NFC', s).lower()
    s = re.sub(r'https?://\S+', ' ', s)
    s = re.sub(r'[^\w\sà-ỹđ]', ' ', s)
    return re.sub(r'\s+', ' ', s).strip()


def page_text(slug: str) -> str:
    p = ROOT / 'dist' / 'bai' / slug / 'index.html'
    if not p.exists():
        return ''
    t = p.read_text('utf-8')
    t = re.sub(r'<script[\s\S]*?</script>|<style[\s\S]*?</style>', ' ', t)
    # giữ chữ trong thuộc tính alt/aria-label/title (ảnh, sơ đồ)
    attrs = ' '.join(re.findall(r'(?:alt|aria-label|title)="([^"]*)"', t))
    t = re.sub(r'<[^>]+>', ' ', t)
    return norm(html.unescape(t + ' ' + attrs))


def sentences(path: Path):
    d = docx.Document(str(path))
    texts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            texts += [c.text for c in r.cells]
    for t in texts:
        for s in re.split(r'(?<=[.!?…:])\s+|\n', t):
            n = norm(s)
            if len(n) >= 25:
                yield s.strip(), n


total_ok = total = 0
rows = []
for path in sorted((ROOT / 'tai-lieu-goc').rglob('*.docx')):
    code = path.name.split(' - ')[0]
    slug = SLUGS.get(code)
    if not slug:
        continue
    page = page_text(slug)
    if not page:
        rows.append((slug, 0, 0, ['(chưa có trang)']))
        continue
    ok, missing, n = 0, [], 0
    for raw, s in sentences(path):
        n += 1
        # khớp cả câu, hoặc 2 nửa câu (câu dài bị tách bởi sơ đồ / định dạng)
        half = len(s) // 2
        if s in page or (len(s) > 60 and s[:half] in page and s[half:] in page):
            ok += 1
        else:
            missing.append(raw)
    total_ok += ok
    total += n
    rows.append((slug, ok, n, missing))

for slug, ok, n, missing in rows:
    pct = 100 * ok / n if n else 0
    flag = 'OK ' if pct >= 97 else '!! '
    print(f'{flag}{slug:<46} {ok:>4}/{n:<4} {pct:5.1f}%')
    if SHOW or pct < 97:
        for m in missing[:12]:
            print('      – ' + m[:140])
print(f'\nTổng: {total_ok}/{total} câu ({100 * total_ok / max(total, 1):.1f}%)')
